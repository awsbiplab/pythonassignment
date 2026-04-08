
from docx import Document

def generate_full_report(df):
    doc = Document()
    doc.add_heading('Hero FinCorp Analysis', 0)
    doc.add_paragraph(f"Total Loans: {len(df)}")
    doc.add_paragraph(f"Default Rate: {df['DEFAULT_FLAG'].mean():.2%}")
    doc.save("reports/hero_fincorp_analysis.docx")
