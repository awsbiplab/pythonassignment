from docx import Document

def generateReport(metrics):
    doc = Document()

    doc.add_heading('Hero FinCorp Analysis', 0)

    doc.add_heading('Executive Summary', 1)
    doc.add_paragraph(f"Default Rate: {metrics['default_rate']:.2%}")

    doc.add_heading('Key Insights', 1)
    doc.add_paragraph("Low credit score customers have higher risk.")

    doc.add_heading('Recommendations', 1)
    doc.add_paragraph("Improve credit checks and reduce processing time.")

    doc.save('reports/hero_fincorp_analysis.docx')